from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\alvar\Documents\GimnasioWapp")
OUT_DIR = ROOT / "docs" / "contrato-cliente"
ASSETS_DIR = ROOT / "docs" / "contract-assets"
OUT_DIR.mkdir(parents=True, exist_ok=True)
ASSETS_DIR.mkdir(parents=True, exist_ok=True)

DOC_TITLE = "Contrato de desarrollo y mantenimiento - Ronquillo Te Cuida"
DEVELOPER = "Alvaro Pavon Martinez"
APP_NAME = "Ronquillo Te Cuida"
DOC_DATE = date(2026, 6, 13).strftime("%d/%m/%Y")

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(30, 95, 160)
MUTED = RGBColor(96, 112, 128)
LIGHT_BLUE = "E8EEF5"
FONT = "Calibri"


def load_font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf") if bold else Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf") if bold else Path(r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def safe_open(path: Path) -> Image.Image:
    return Image.open(path).convert("RGB")


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None, size: float = 9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color="D7DBE2", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def make_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_cm: Sequence[float] | None = None,
    font_size=8,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    if widths_cm:
        for i, width in enumerate(widths_cm):
            for cell in table.columns[i].cells:
                cell.width = Cm(width)
    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=NAVY, size=font_size)
    repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=font_size)
            if widths_cm:
                cells[i].width = Cm(widths_cm[i])
    set_table_borders(table)
    doc.add_paragraph()
    return table


def add_caption(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.color.rgb = BLUE if level < 3 else NAVY
    return paragraph


def add_body(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(10.5)
    return paragraph


def add_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.name = FONT
        run.font.size = Pt(10.2)


def add_numbered(doc: Document, items: Iterable[str]):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.name = FONT
        run.font.size = Pt(10.2)


def add_callout(doc: Document, title: str, body: str, fill: str = "FFF8E8"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="E2C879")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = NAVY
    run.font.size = Pt(10.5)
    paragraph2 = cell.add_paragraph()
    paragraph2.paragraph_format.space_after = Pt(0)
    run2 = paragraph2.add_run(body)
    run2.font.size = Pt(10)
    doc.add_paragraph()


def image_card(img: Image.Image, target_w: int, target_h: int, label: str) -> Image.Image:
    canvas = Image.new("RGB", (target_w, target_h), "white")
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle([0, 0, target_w - 1, target_h - 1], radius=18, outline=(209, 215, 224), width=2, fill=(255, 255, 255))
    draw.text((20, 14), label, fill=(11, 37, 69), font=load_font(24, True))
    title_h = 54
    content = ImageOps.contain(img, (target_w - 28, target_h - title_h - 22))
    x = (target_w - content.width) // 2
    y = title_h + (target_h - title_h - content.height) // 2
    canvas.paste(content, (x, y))
    return canvas


def make_screenshot_collages():
    web_paths = [
        ASSETS_DIR / "web-login.png",
        ASSETS_DIR / "web-dashboard-admin.png",
        ASSETS_DIR / "web-users-admin.png",
        ASSETS_DIR / "web-payments-admin.png",
    ]
    web_labels = ["Login web", "Panel admin", "Directorio", "Pagos y avisos"]
    cards = [image_card(safe_open(path), 920, 560, label) for path, label in zip(web_paths, web_labels) if path.exists()]
    if cards:
        collage = Image.new("RGB", (1880, 1160), (245, 247, 250))
        for i, card in enumerate(cards[:4]):
            collage.paste(card, (20 + (i % 2) * 930, 20 + (i // 2) * 570))
        collage.save(ASSETS_DIR / "screens-web.jpg", quality=88)

    ios_dir = Path(r"C:\Users\alvar\Desktop\Ronquillo Te Cuida - iOS App Store\02-capturas-app-store-1284x2778")
    android_dir = Path(r"C:\Users\alvar\Desktop\Ronquillo Te Cuida - Android Play Store")
    mobile_paths = [
        ios_dir / "ronquillo-ios-screenshot-01-1284x2778.png",
        ios_dir / "ronquillo-ios-screenshot-04-1284x2778.png",
        android_dir / "Screenshot_20260603_233402_Ronquillo Te Cuida.jpg",
        android_dir / "Screenshot_20260603_233416_Ronquillo Te Cuida.jpg",
    ]
    mobile_labels = ["iOS - Inicio", "iOS - Gestión", "Android - Inicio", "Android - Reservas"]
    mobile_cards = [image_card(safe_open(path), 420, 860, label) for path, label in zip(mobile_paths, mobile_labels) if path.exists()]
    if mobile_cards:
        collage = Image.new("RGB", (1780, 910), (245, 247, 250))
        for i, card in enumerate(mobile_cards[:4]):
            collage.paste(card, (20 + i * 440, 25))
        collage.save(ASSETS_DIR / "screens-mobile.jpg", quality=88)


def draw_box(draw, xy, title, lines, fill=(255, 255, 255), outline=(35, 75, 120)):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    draw.text((x1 + 18, y1 + 14), title, fill=(11, 37, 69), font=load_font(26, True))
    y = y1 + 54
    for line in lines:
        draw.text((x1 + 18, y), line, fill=(40, 50, 65), font=load_font(19))
        y += 27


def arrow(draw, start, end, color=(30, 95, 160), width=4):
    import math

    draw.line([start, end], fill=color, width=width)
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 14
    pts = [
        end,
        (end[0] - size * math.cos(ang - 0.45), end[1] - size * math.sin(ang - 0.45)),
        (end[0] - size * math.cos(ang + 0.45), end[1] - size * math.sin(ang + 0.45)),
    ]
    draw.polygon(pts, fill=color)


def create_architecture_diagram():
    img = Image.new("RGB", (1800, 1050), (245, 247, 250))
    d = ImageDraw.Draw(img)
    d.text((60, 40), "Arquitectura general del sistema", fill=(11, 37, 69), font=load_font(42, True))
    draw_box(d, (70, 150, 420, 340), "Clientes", ["Web React/Vite", "Android Expo", "iOS Expo"])
    draw_box(d, (560, 150, 930, 340), "Nginx + HTTPS", ["Dominio DuckDNS", "Proxy inverso", "SSL / endpoint público"], fill=(255, 253, 245), outline=(199, 143, 35))
    draw_box(d, (1070, 130, 1560, 370), "API Node.js + Express", ["JWT y roles", "Validación Zod", "Prisma ORM", "WebSocket tiempo real"], fill=(238, 246, 255))
    draw_box(d, (1160, 500, 1540, 705), "MySQL 8 / InnoDB", ["Usuarios y roles", "Clases y reservas", "Pagos y avisos", "Penalizaciones"])
    draw_box(d, (560, 500, 930, 705), "Servicios background", ["node-cron", "No-show < 30 min", "Recordatorio 45 min", "Vencimiento cuotas"])
    draw_box(d, (70, 500, 420, 705), "Notificaciones push", ["expo-server-sdk", "expo-notifications", "APNs / FCM", "Lista de espera"])
    for start, end in [((420, 245), (560, 245)), ((930, 245), (1070, 245)), ((1310, 370), (1340, 500)), ((1070, 315), (930, 570)), ((560, 600), (420, 600)), ((1160, 610), (930, 610))]:
        arrow(d, start, end)
    d.text((70, 820), "Principio clave: una única API gobierna web, Android e iOS. La base de datos se bloquea transaccionalmente para evitar overbooking.", fill=(64, 75, 90), font=load_font(24, True))
    img.save(ASSETS_DIR / "diagram-architecture.png", quality=95)


def create_flow_diagram():
    img = Image.new("RGB", (1800, 1150), (245, 247, 250))
    d = ImageDraw.Draw(img)
    d.text((60, 40), "Flujo crítico de reserva, asistencia y lista de espera", fill=(11, 37, 69), font=load_font(40, True))
    boxes = [
        ((70, 150, 390, 300), "Cliente", ["Inicia sesión", "Estado cuota PAGADO"]),
        ((520, 150, 850, 300), "Reserva", ["API valida rol", "Valida cuota activa"]),
        ((980, 150, 1320, 300), "Bloqueo DB", ["SELECT ... FOR UPDATE", "Cuenta plazas reales"]),
        ((1450, 150, 1740, 300), "Resultado", ["CONFIRMADA", "o EN_ESPERA"]),
        ((70, 480, 390, 650), "Validación", ["Hasta 30 min antes", "Cliente confirma asistencia"]),
        ((520, 480, 850, 650), "Cron no-show", ["< 30 min", "CONFIRMADA -> NO_ASISTE", "Crea penalización"]),
        ((980, 480, 1320, 650), "Promoción", ["No penalizados primero", "Luego fecha_solicitud"]),
        ((1450, 480, 1740, 650), "Push", ["Plaza confirmada", "Sin recargar app"]),
        ((520, 830, 850, 1000), "Pago cuota", ["Cliente reporta pago", "Admin recibe aviso"]),
        ((980, 830, 1320, 1000), "Renovación", ["Admin confirma", "Nueva fecha vencimiento"]),
    ]
    for xy, title, lines in boxes:
        draw_box(d, xy, title, lines)
    for start, end in [((390, 225), (520, 225)), ((850, 225), (980, 225)), ((1320, 225), (1450, 225)), ((230, 300), (230, 480)), ((690, 300), (690, 480)), ((850, 565), (980, 565)), ((1320, 565), (1450, 565)), ((690, 650), (690, 830)), ((850, 915), (980, 915))]:
        arrow(d, start, end)
    d.text((70, 1070), "Regla de negocio: profesores y administradores están exentos de cuota; solo clientes necesitan membresía activa para reservar.", fill=(64, 75, 90), font=load_font(23, True))
    img.save(ASSETS_DIR / "diagram-flow.png", quality=95)


def create_er_diagram():
    img = Image.new("RGB", (2000, 1450), "white")
    d = ImageDraw.Draw(img)
    d.text((55, 35), "Diagrama entidad-relación lógico", fill=(11, 37, 69), font=load_font(42, True))
    entities = {
        "Usuarios": ((70, 140, 445, 390), ["id PK", "nombre", "email UNIQUE", "password_hash", "rol", "estado_mensualidad", "membership_expires_at"]),
        "DispositivosPush": ((70, 520, 445, 720), ["id PK", "user_id FK", "push_token UNIQUE", "plataforma"]),
        "PagosMensualidad": ((70, 850, 445, 1120), ["id PK", "user_id FK", "amount_cents", "status", "paid_at", "reviewed_by_id FK"]),
        "TiposClase": ((760, 140, 1110, 300), ["id PK", "nombre UNIQUE"]),
        "Clases": ((760, 420, 1110, 700), ["id PK", "titulo", "tipo_clase_id FK", "teacher_id FK", "capacidad_maxima", "fecha_hora_inicio", "fecha_hora_fin"]),
        "Reservas": ((1380, 420, 1760, 700), ["id PK", "user_id FK", "clase_id FK", "estado", "fecha_solicitud", "UNIQUE(user, clase)"]),
        "Penalizaciones": ((1380, 830, 1760, 1060), ["id PK", "user_id FK", "tipo_clase_id FK", "activa"]),
        "NotificacionesAdmin": ((760, 850, 1110, 1120), ["id PK", "type", "title", "message", "user_id FK", "payment_id FK", "read_at"]),
        "ImageBank": ((760, 1160, 1110, 1310), ["id PK", "keyword UNIQUE", "image_url"]),
        "SystemSettings": ((1380, 1160, 1760, 1310), ["id PK", "app_name", "hero_image", "updated_at"]),
    }
    for name, (xy, fields) in entities.items():
        x1, y1, x2, y2 = xy
        d.rounded_rectangle(xy, radius=16, fill=(248, 250, 252), outline=(30, 95, 160), width=3)
        d.rectangle([x1, y1, x2, y1 + 42], fill=(232, 238, 245), outline=(30, 95, 160), width=2)
        d.text((x1 + 14, y1 + 9), name, fill=(11, 37, 69), font=load_font(23, True))
        y = y1 + 55
        for field in fields:
            d.text((x1 + 16, y), field, fill=(40, 50, 65), font=load_font(19))
            y += 28

    def mid_right(name):
        x1, y1, x2, y2 = entities[name][0]
        return x2, (y1 + y2) // 2

    def mid_left(name):
        x1, y1, x2, y2 = entities[name][0]
        return x1, (y1 + y2) // 2

    def mid_bottom(name):
        x1, y1, x2, y2 = entities[name][0]
        return (x1 + x2) // 2, y2

    def mid_top(name):
        x1, y1, x2, y2 = entities[name][0]
        return (x1 + x2) // 2, y1

    for start, end in [
        (mid_bottom("Usuarios"), mid_top("DispositivosPush")),
        (mid_bottom("DispositivosPush"), mid_top("PagosMensualidad")),
        (mid_right("Usuarios"), mid_left("Clases")),
        (mid_bottom("TiposClase"), mid_top("Clases")),
        (mid_right("Clases"), mid_left("Reservas")),
        (mid_right("Usuarios"), mid_left("Reservas")),
        (mid_bottom("Reservas"), mid_top("Penalizaciones")),
        (mid_right("TiposClase"), mid_left("Penalizaciones")),
        (mid_right("PagosMensualidad"), mid_left("NotificacionesAdmin")),
    ]:
        arrow(d, start, end, color=(120, 130, 145), width=3)
    d.text((60, 1375), "Relaciones principales: Usuario 1-N Reservas, Clase 1-N Reservas, TipoClase 1-N Clases/Penalizaciones, Usuario 1-N Pagos/Dispositivos.", fill=(64, 75, 90), font=load_font(22, True))
    img.save(ASSETS_DIR / "diagram-er.png", quality=95)


def generate_diagrams_and_collages():
    make_screenshot_collages()
    create_architecture_diagram()
    create_flow_diagram()
    create_er_diagram()


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, NAVY)]:
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{DOC_TITLE} | {DEVELOPER} | Página ")
    run.font.name = FONT
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    add_page_number(paragraph)


def add_cover(doc: Document):
    logo = ROOT / "mobile" / "assets" / "logo.jpg"
    if logo.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(logo), width=Cm(7.0))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    run = paragraph.add_run("Contrato de desarrollo, publicación y mantenimiento")
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(24)
    run.font.color.rgb = NAVY

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(APP_NAME)
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(30)
    run.font.color.rgb = BLUE

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Sistema integral web, Android e iOS para gestión de clases, reservas, cuotas y notificaciones")
    run.font.name = FONT
    run.font.size = Pt(13)
    run.font.color.rgb = MUTED

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    data = [
        ("Proveedor / desarrollador", DEVELOPER),
        ("Cliente", "Ronquillo Te Cuida / Titular del centro"),
        ("Fecha del documento", DOC_DATE),
        ("Importe de entrega ofertado", "500 € - precio especial con costes iniciales incluidos"),
        ("Mantenimiento", "50 €/mes - precio especial mientras el servicio siga activo"),
    ]
    for row, (label, value) in zip(table.rows, data):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[0], label, bold=True, color=NAVY, size=9.5)
        set_cell_text(row.cells[1], value, size=9.5)
    doc.add_paragraph()
    add_callout(doc, "Resumen comercial", "El valor real de mercado de una plataforma de este alcance es muy superior. Se aplica una oferta cerrada de 500 € para la entrega y 50 €/mes para mantenimiento, licencias, servidores y soporte evolutivo ordinario.")
    doc.add_page_break()


def add_index(doc: Document):
    add_heading(doc, "Índice", 1)
    add_numbered(doc, [
        "Objeto y alcance del contrato",
        "Entregables realizados",
        "Funcionalidades por rol",
        "Arquitectura técnica y tecnologías",
        "Modelo de datos y flujos críticos",
        "Seguridad, validación y calidad",
        "Costes de producción y oferta aplicada",
        "Mantenimiento mensual y condiciones de baja",
        "Condiciones contractuales",
        "Evidencia visual",
        "Anexos técnicos y firmas",
    ])
    doc.add_page_break()


def add_scope(doc: Document):
    add_heading(doc, "1. Objeto y alcance del contrato", 1)
    add_body(doc, "El presente documento formaliza la entrega técnica, funcional y comercial del sistema Ronquillo Te Cuida, compuesto por una plataforma web, una aplicación Android, una aplicación iOS y un backend centralizado con base de datos, autenticación, gestión de reservas, control de cuotas, notificaciones push y sincronización en tiempo real.")
    add_body(doc, "El sistema permite operar el flujo completo de un centro deportivo: alta y gestión de usuarios, creación de clases, asignación de profesores, reservas de clientes, listas de espera, validación de asistencia, penalizaciones por no asistencia, control de mensualidades, avisos administrativos, personalización de marca y publicación en tiendas móviles.")
    add_callout(doc, "Estado de publicación", "La aplicación iOS figura como aprobada y publicada según confirmación del titular de la cuenta Apple Developer. Android queda preparado, firmado y empaquetado para Play Console con el identificador com.azrael.ronquillotecuida; su distribución final depende de los requisitos y plazos de Google Play.", fill="E8F5E9")
    add_heading(doc, "1.1 Partes", 2)
    make_table(doc, ["Parte", "Identificación", "Responsabilidad"], [
        ["Desarrollador", DEVELOPER, "Diseño, desarrollo, despliegue, administración técnica, publicación y mantenimiento según este contrato."],
        ["Cliente", "Titular de Ronquillo Te Cuida", "Uso del sistema, comunicación de necesidades operativas, validación de pagos/cuotas y aportación de datos/contenidos correctos."],
    ], [3.2, 5.0, 8.8], font_size=8.5)
    add_heading(doc, "1.2 Plataforma entregada", 2)
    add_bullets(doc, [
        "Backend API con Node.js, Express, Prisma y MySQL, desplegado tras Nginx y accesible por HTTPS.",
        "Panel web responsive en React/Vite/Tailwind para administración completa del centro.",
        "Aplicación móvil Android e iOS con Expo/React Native, notificaciones push y sesión recordada.",
        "Sistema transaccional de reservas, lista de espera y penalizaciones diseñado para evitar overbooking.",
        "Sistema de cuotas con fecha de vencimiento, exención automática para administradores/profesores y avisos al administrador.",
        "Sincronización en tiempo real para web y móvil mediante canal WebSocket/eventos de cambio.",
    ])


def add_deliverables(doc: Document):
    add_heading(doc, "2. Entregables realizados", 1)
    make_table(doc, ["Área", "Entregado", "Detalle"], [
        ["Aplicación web", "Sí", "Login, dashboard por rol, gestión de clases, usuarios, pagos, imágenes, métricas, perfil, marca y política de privacidad."],
        ["Aplicación Android", "Sí", "App Expo/React Native con paquete com.azrael.ronquillotecuida, AAB firmado, notificaciones y funcionalidades equivalentes adaptadas a móvil."],
        ["Aplicación iOS", "Sí", "Bundle com.azrael.ronquillotecuida, publicación gestionada en App Store Connect, APNs configurado y soporte iPhone."],
        ["Backend/API", "Sí", "Express 5, JWT, roles, Zod, Prisma, MySQL, cron jobs, WebSocket y endpoints REST documentados."],
        ["Base de datos", "Sí", "Modelo MySQL/InnoDB con usuarios, dispositivos push, clases, tipos, reservas, penalizaciones, pagos, avisos y ajustes."],
        ["Despliegue", "Sí", "Servidor Linux con Nginx, dominio ronquillotecuida.duckdns.org, HTTPS y política de privacidad pública."],
        ["Pruebas", "Sí", "Jest/Supertest para reservas, overbooking, lista de espera, penalizaciones, cron jobs, membresías, usuarios y tiempo real."],
        ["Documentación", "Sí", "README/API, política de privacidad, preparación de tiendas y este contrato con anexo técnico."],
    ], [3.1, 2.0, 11.9])
    add_heading(doc, "2.1 Publicación y distribución", 2)
    add_bullets(doc, [
        "Nombre comercial: Ronquillo Te Cuida.",
        "Dominio y API pública: https://ronquillotecuida.duckdns.org.",
        "URL de política de privacidad: https://ronquillotecuida.duckdns.org/privacy-policy.html.",
        "Identificador Android: com.azrael.ronquillotecuida.",
        "Identificador iOS: com.azrael.ronquillotecuida.",
        "App Store Connect ID usado durante el proceso: 6778762657.",
        "Versión comercial inicial: 1.0.0.",
    ])


def add_roles(doc: Document):
    add_heading(doc, "3. Funcionalidades por rol", 1)
    make_table(doc, ["Rol", "Funcionalidades incluidas"], [
        ["Administrador", "Crear/editar/eliminar usuarios; cambiar contraseñas; gestionar roles; renovar cuotas; revisar pagos; crear/editar/eliminar clases; asignar profesores; gestionar banco de imágenes; personalizar marca; consultar métricas; ver avisos; eliminar usuarios de clases."],
        ["Profesor", "Ver sus clases; crear o editar clases propias cuando procede; consultar detalle de clase; gestionar asistencia/reservas de sus clases; eliminar alumnos de sus clases."],
        ["Cliente/Alumno", "Iniciar sesión; consultar clases; reservar; entrar en lista de espera; cancelar reserva; validar asistencia hasta 30 minutos antes; reportar pago; recibir notificaciones; ver estado de cuota y vencimiento."],
    ], [3.0, 14.0], font_size=8.5)
    add_heading(doc, "3.1 Reglas de negocio implementadas", 2)
    add_bullets(doc, [
        "Solo los clientes con cuota PAGADO y fecha de vencimiento futura pueden reservar clases.",
        "Administradores y profesores están exentos de pago; el backend fuerza su estado como exento/PAGADO sin vencimiento.",
        "Las clases llenas mandan nuevas reservas a EN_ESPERA, no a confirmación, evitando sobreaforo.",
        "Si se libera plaza, se promociona a la lista de espera priorizando usuarios sin penalización activa para el tipo de clase y después por fecha de solicitud.",
        "El cliente debe validar asistencia hasta 30 minutos antes; si no lo hace, el cron lo marca como NO_ASISTE, crea penalización y libera plaza.",
        "Si un usuario con penalización activa asiste correctamente a una clase del mismo tipo, la penalización queda desactivada.",
        "A los 45 minutos antes de la clase se envían recordatorios push a reservas confirmadas pendientes de validar.",
        "El pago de cuota puede ser reportado por el cliente y confirmado por el administrador, renovando automáticamente la fecha de vencimiento.",
    ])


def add_architecture(doc: Document):
    add_heading(doc, "4. Arquitectura técnica y tecnologías", 1)
    architecture = ASSETS_DIR / "diagram-architecture.png"
    if architecture.exists():
        doc.add_picture(str(architecture), width=Cm(16.5))
        add_caption(doc, "Figura 1. Arquitectura general de la solución entregada.")
    add_heading(doc, "4.1 Versiones y stack tecnológico", 2)
    make_table(doc, ["Capa", "Tecnología", "Versión/criterio verificado", "Uso"], [
        ["Runtime", "Node.js", "20.20.1 / engine >=20.19", "Backend, tooling y scripts de build."],
        ["Backend", "Express", "5.2.1", "API REST, middleware, seguridad y rutas."],
        ["ORM", "Prisma", "6.19.3", "Acceso tipado a MySQL y migraciones."],
        ["Base de datos", "MySQL", "8.4 Docker / compatible 8.0+", "Persistencia InnoDB de usuarios, reservas, clases, pagos y avisos."],
        ["Validación", "Zod", "4.4.3", "Validación estricta de payloads de API."],
        ["Autenticación", "JWT jsonwebtoken", "9.0.3", "Sesión segura por token y roles."],
        ["Hash de contraseñas", "bcryptjs", "3.0.3", "Hash de contraseñas antes de guardar."],
        ["Cron", "node-cron", "4.2.1", "Recordatorios, no-show y vencimiento de cuotas."],
        ["Push backend", "expo-server-sdk", "6.1.0", "Envío de notificaciones push a Expo/APNs/FCM."],
        ["Tiempo real", "ws", "8.21.0", "Sincronización inmediata web/móvil."],
        ["Web", "React", "19.2.3", "Interfaz web por roles."],
        ["Web tooling", "Vite", "8.0.14", "Build moderno del frontend."],
        ["Estilos web", "TailwindCSS", "4.3.0", "UI responsive y pulida."],
        ["HTTP cliente", "Axios", "1.16.1", "Consumo API en web y móvil."],
        ["Mobile", "Expo", "~56.0.9", "Build y publicación Android/iOS."],
        ["Mobile", "React Native", "0.85.3", "Interfaz nativa multiplataforma."],
        ["Mobile", "React", "19.2.3", "Componentes y estado UI."],
        ["Mobile storage", "AsyncStorage", "2.2.0", "Recordar sesión del usuario."],
        ["Notificaciones móvil", "expo-notifications", "~56.0.16", "Registro y recepción de push."],
        ["Testing", "Jest + Supertest", "30.4.2 / 7.2.2", "Pruebas unitarias e integración API."],
        ["Lenguaje", "TypeScript", "5.9.3 API/Web, ~6.0.3 móvil", "Tipado y control de build."],
        ["Contenedores", "Docker", "29.2.1 verificado local", "MySQL y soporte de despliegue."],
    ], [2.5, 3.2, 3.0, 8.0], font_size=7.6)
    add_heading(doc, "4.2 Arquitectura de carpetas", 2)
    add_bullets(doc, [
        "Monorepo con workspaces: api, web y mobile.",
        "api: Express, Prisma, rutas, servicios, middleware, cron jobs, tests y documentación API.",
        "web: React/Vite con componentes de administración, cliente, profesor, login y cliente API/realtime.",
        "mobile: Expo/React Native con pantallas de login/dashboard, cliente API, realtime y registro push.",
    ])


def add_data_and_flows(doc: Document):
    add_heading(doc, "5. Modelo de datos y flujos críticos", 1)
    er = ASSETS_DIR / "diagram-er.png"
    flow = ASSETS_DIR / "diagram-flow.png"
    if er.exists():
        doc.add_picture(str(er), width=Cm(17.0))
        add_caption(doc, "Figura 2. Modelo entidad-relación lógico de la base de datos.")
    if flow.exists():
        doc.add_picture(str(flow), width=Cm(17.0))
        add_caption(doc, "Figura 3. Flujo transaccional de reserva, asistencia, penalización y promoción de lista de espera.")
    add_heading(doc, "5.1 Entidades principales", 2)
    make_table(doc, ["Entidad", "Propósito"], [
        ["Usuarios", "Identidad, rol, email, contraseña cifrada, cuota, vencimiento, teléfono y foto."],
        ["DispositivosPush", "Tokens Expo asociados a usuario y plataforma para enviar avisos nativos."],
        ["TiposClase", "Catálogo de tipos como Yoga, Spinning o Functional."],
        ["Clases", "Sesiones con título, descripción, profesor, tipo, capacidad, fechas e imagen."],
        ["Reservas", "Solicitud del cliente, estado de reserva, lista de espera, asistencia, no asistencia o cancelación."],
        ["Penalizaciones", "Restricción activa por usuario y tipo de clase cuando existe no asistencia."],
        ["PagosMensualidad", "Pagos reportados por clientes y revisados por administradores."],
        ["NotificacionesAdmin", "Avisos internos de pagos reportados y cuotas vencidas."],
        ["ImageBank/SystemSettings", "Personalización visual de clases y marca de la aplicación."],
    ], [4.0, 13.0], font_size=8.4)


def add_security_quality(doc: Document):
    add_heading(doc, "6. Seguridad, validación y calidad", 1)
    add_heading(doc, "6.1 Seguridad aplicada", 2)
    add_bullets(doc, [
        "Autenticación JWT para todas las rutas privadas.",
        "Control de roles ADMIN, TEACHER y CLIENT en middleware de backend.",
        "Contraseñas almacenadas como hash bcrypt; nunca como texto plano.",
        "Validación de entradas con Zod en rutas críticas.",
        "Uso de Prisma ORM para evitar inyección SQL en operaciones normales.",
        "Consultas raw críticas realizadas con templates parametrizados de Prisma; sin concatenar input de usuario.",
        "Helmet, CORS controlado y rate limiting en Express.",
        "HTTPS en producción detrás de Nginx.",
        "Política de privacidad pública y enlaces desde web/móvil.",
        "No se incluyen credenciales sensibles en este contrato.",
    ])
    add_heading(doc, "6.2 Calidad y pruebas", 2)
    make_table(doc, ["Tipo de prueba", "Cobertura"], [
        ["Reservas", "Prevención de overbooking, confirmación/lista de espera, cancelación y promoción."],
        ["Lista de espera", "Prioridad de usuarios sin penalización activa y orden por fecha de solicitud."],
        ["Cron jobs", "Cambio a NO_ASISTE, creación de penalizaciones, liberación de plaza y recordatorios."],
        ["Membresías", "Reporte de pago, confirmación admin, renovación, vencimiento y bloqueo de reservas."],
        ["Usuarios", "Alta, edición, cambio de contraseña por administrador y roles exentos."],
        ["Tiempo real", "Broadcast de cambios para evitar recargas manuales."],
        ["Móvil", "Validación en dispositivo Android físico y build iOS con credenciales de App Store."],
    ], [4.0, 13.0], font_size=8.4)
    add_callout(doc, "Punto técnico relevante", "La prevención de sobreaforo no depende solo de la interfaz. El backend bloquea filas de clase/reservas con SELECT ... FOR UPDATE dentro de una transacción, que es la pieza correcta para concurrencia real.", fill="E8EEF5")


def add_costs(doc: Document):
    add_heading(doc, "7. Costes de producción y oferta aplicada", 1)
    add_body(doc, "Para contextualizar el precio final, se han revisado referencias públicas de mercado de 2026 sobre desarrollo web, móvil y mantenimiento. Las cifras de referencia se usan solo como comparación comercial; el importe pactado para este cliente es una oferta especial cerrada.")
    make_table(doc, ["Referencia de mercado", "Rango publicado", "Aplicación al proyecto"], [
        ["Aplicación móvil básica/medio alcance", "Appinventiv: desde 40.000 USD para apps básicas y 100.000-200.000 USD para complejidad moderada.", "Ronquillo Te Cuida incluye Android+iOS, backend, autenticación, reservas, push y roles."],
        ["Aplicación web", "SaM Solutions: 20.000-70.000 USD para web apps simples y 80.000-180.000 USD para complejidad media.", "La web entregada es un panel operativo con roles, dashboards, CRUD, pagos y sincronización."],
        ["Web app pequeña de negocio", "Digisoft: 15.000-50.000 USD para small business web app; mantenimiento 250-1.500 USD/mes.", "Sirve como referencia conservadora solo para la parte web."],
        ["Mantenimiento móvil", "AppsChopper/Appinventiv: 15-25% o 15-20% anual del coste inicial.", "Con un coste de mercado bajo de 15.000 €, el mantenimiento anual típico ya sería 2.250-3.750 €."],
        ["Mantenimiento mensual", "Imaginovation: 2.500-5.000 USD/mes como benchmark inicial; otros informes sitúan apps simples en miles/año.", "La cuota ofrecida de 50 €/mes es muy inferior a referencias de mercado."],
        ["Licencias de tienda", "Google Play: 25 USD pago único. Apple Developer Program: 99 USD/año.", "Costes iniciales incluidos en la oferta; renovaciones anuales cubiertas por mantenimiento activo."],
    ], [4.0, 4.7, 8.3], font_size=7.8)
    add_heading(doc, "7.1 Valor estimado conservador", 2)
    make_table(doc, ["Concepto", "Referencia conservadora", "Observación"], [
        ["Backend/API + base de datos", "3.000-8.000 €", "Autenticación, roles, reservas, cron, push, tiempo real y despliegue."],
        ["Panel web", "5.000-15.000 €", "Panel administrativo con gestión completa del centro."],
        ["Android + iOS", "8.000-25.000 €", "Apps móviles con mismo backend, notificaciones y publicación."],
        ["Testing, documentación y publicación", "1.500-5.000 €", "Pruebas, ajustes de tienda, capturas, política de privacidad y builds."],
        ["Total conservador", "17.500-53.000 €", "Estimación deliberadamente inferior a muchas referencias internacionales."],
    ], [5.0, 4.0, 8.0], font_size=8.2)
    add_callout(doc, "Oferta final al cliente", "Se deja el proyecto completo en 500 €, incluyendo la entrega inicial, configuración, publicación, costes iniciales de licencias y puesta en marcha. Es un precio especial y muy inferior al valor de mercado estimado.")
    add_heading(doc, "7.2 Fuentes de referencia", 2)
    add_bullets(doc, [
        "Appinventiv - Mobile App Development Cost 2026: https://appinventiv.com/guide/mobile-app-development-cost/",
        "SaM Solutions - Web App Development Cost 2026: https://sam-solutions.com/blog/web-app-development-cost/",
        "Digisoft Solution - Web App Development Cost Summary 2026: https://www.digisoftsolution.com/blog/web-app-development-cost",
        "Appinventiv - Mobile App Maintenance Costs 2026: https://appinventiv.com/blog/what-is-the-cost-to-maintain-an-app/",
        "AppsChopper - App Maintenance Costs 2026: https://www.appschopper.com/blog/cost-to-maintain-an-app/",
        "Imaginovation - Mobile App Maintenance Cost 2026: https://imaginovation.net/blog/importance-mobile-app-maintenance-cost/",
        "Google Play Console Help - registration fee: https://support.google.com/googleplay/android-developer/answer/6112435",
        "Apple Developer Support - membership: https://developer.apple.com/support/compare-memberships/",
    ])


def add_maintenance(doc: Document):
    add_heading(doc, "8. Mantenimiento mensual y condiciones de baja", 1)
    add_body(doc, "El mantenimiento contratado tiene un precio especial de 50 €/mes. Esta cuota se mantiene mientras el alcance técnico y el volumen de uso sigan dentro de las condiciones actuales del proyecto.")
    make_table(doc, ["Incluido en 50 €/mes", "Descripción"], [
        ["Administración de código", "Custodia técnica del repositorio, control de versiones, preparación de builds y gestión de cambios."],
        ["Actualizaciones críticas", "Parches de seguridad, dependencias críticas, ajustes por cambios obligatorios de tiendas o sistema operativo."],
        ["Cambios solicitados por el cliente", "Cambios razonables de evolución ordinaria, textos, ajustes de interfaz, campos, pequeños flujos y mejoras operativas previamente acordadas."],
        ["Licencias y renovaciones", "Renovaciones anuales necesarias para mantener la app disponible, incluyendo costes de tiendas cuando correspondan."],
        ["Servidor y despliegue", "Administración básica del servidor, Nginx, dominio/subdominio utilizado, certificados, despliegues y reinicios controlados."],
        ["Soporte funcional", "Revisión de incidencias reportadas, acompañamiento de uso y pequeñas correcciones."],
    ], [5.0, 12.0], font_size=8.2)
    add_heading(doc, "8.1 No incluido salvo acuerdo adicional", 2)
    add_bullets(doc, [
        "Rediseño completo de marca o interfaz desde cero.",
        "Nuevas plataformas distintas de web, Android e iOS.",
        "Integraciones de pago externas no previstas, pasarelas bancarias, facturación electrónica compleja o sistemas de terceros de gran alcance.",
        "Aumento extraordinario de infraestructura por crecimiento significativo de usuarios/tráfico.",
        "Cesión, exportación o traspaso de código, repositorios, credenciales, servidores, tiendas o infraestructura.",
    ])
    add_heading(doc, "8.2 Continuidad del servicio y suspensión por baja", 2)
    add_body(doc, "El mantenimiento permanecerá activo mientras el cliente mantenga al día la cuota mensual pactada. La continuidad operativa de la web, API, base de datos, certificados, builds móviles, licencias y servidores queda vinculada a dicho mantenimiento, porque el sistema requiere administración técnica continua para seguir funcionando de forma segura.")
    add_body(doc, "Si el cliente solicita la baja del mantenimiento, deja de abonar la cuota mensual o incumple las condiciones de uso, el desarrollador podrá suspender, apagar o no renovar servidores, certificados, licencias, despliegues, dominios, integraciones y servicios vinculados. En ese caso, la web y las aplicaciones móviles podrán dejar de funcionar total o parcialmente.")
    add_body(doc, "La baja del mantenimiento no implica entrega del código fuente, cesión de repositorios, transferencia de credenciales técnicas, cesión de claves de firma, entrega de scripts internos, traspaso de infraestructura ni obligación del desarrollador de facilitar que un tercero continúe el mantenimiento. La administración técnica queda reservada al desarrollador mientras el servicio esté activo.")


def add_contract_terms(doc: Document):
    add_heading(doc, "9. Condiciones contractuales", 1)
    add_heading(doc, "9.1 Precio y forma de pago", 2)
    add_bullets(doc, [
        "Entrega inicial del proyecto completo: 500 €.",
        "Mantenimiento mensual: 50 €/mes.",
        "Los 500 € incluyen los costes iniciales de licencias, configuración, despliegue y publicación realizados para esta entrega.",
        "Las renovaciones anuales de licencias y servicios se cubrirán con la cuota de mantenimiento mientras esta siga activa.",
        "Si existieran impuestos aplicables no contemplados en este documento, se indicarán en factura o recibo correspondiente según proceda.",
    ])
    add_heading(doc, "9.2 Aceptación del proyecto", 2)
    add_body(doc, "La firma de este documento implica aceptación de la entrega funcional descrita. Cualquier incidencia crítica detectada dentro del uso normal quedará cubierta por mantenimiento siempre que el servicio mensual esté activo.")
    add_heading(doc, "9.3 Propiedad, uso y administración técnica", 2)
    add_body(doc, "El cliente recibe el derecho de uso y explotación operativa de Ronquillo Te Cuida para su centro. El desarrollador mantiene la administración técnica del código, builds, despliegues y configuración mientras exista mantenimiento activo, salvo pacto escrito distinto.")
    add_body(doc, "El cliente no deberá compartir credenciales técnicas, claves de tienda, secretos de servidor, tokens o accesos administrativos con terceros sin coordinación previa, para evitar pérdidas de seguridad o interrupciones del servicio.")
    add_heading(doc, "9.4 Protección de datos y confidencialidad", 2)
    add_body(doc, "El sistema trata datos personales necesarios para su funcionamiento: nombre, email, rol, estado de cuota, reservas, asistencia, penalizaciones, token push y foto si se carga. El cliente será responsable de informar a sus usuarios y usar la plataforma conforme a la normativa aplicable de protección de datos.")
    add_body(doc, "El desarrollador se compromete a no divulgar datos técnicos o personales a terceros salvo obligación legal, necesidad operativa de mantenimiento o autorización expresa del cliente.")
    add_heading(doc, "9.5 Limitación operativa", 2)
    add_body(doc, "La disponibilidad de las aplicaciones publicadas depende también de Apple, Google, Expo, proveedores de hosting, certificados, conectividad del usuario y políticas de tienda. El mantenimiento incluye la gestión razonable de estos elementos, pero no garantiza que terceros externos no cambien requisitos, precios o plazos de revisión.")

    add_heading(doc, "9.6 Propiedad intelectual y código fuente", 2)
    add_body(doc, "El cliente adquiere un derecho de uso operativo de la plataforma Ronquillo Te Cuida para la gestión de su centro, pero no adquiere la propiedad del código fuente, arquitectura interna, scripts, configuraciones técnicas, librerías de integración, credenciales de firma, procedimientos de despliegue ni repositorios utilizados para crear, publicar o mantener la aplicación.")
    add_body(doc, "La titularidad, autoría, derechos de explotación técnica y derechos de reutilización del código desarrollado, salvo contenidos o datos propios aportados por el cliente, quedan reservados al desarrollador Alvaro Pavon Martinez. El precio especial de 500 € corresponde a la entrega funcional y puesta en marcha del sistema, no a una compraventa o cesión completa del código fuente.")
    add_body(doc, "El cliente no podrá exigir la entrega del código fuente, copiarlo, revenderlo, cederlo, publicarlo, entregarlo a terceros, permitir ingeniería inversa, crear obras derivadas no autorizadas ni contratar a otro desarrollador para intervenir sobre el código sin autorización expresa y por escrito del desarrollador.")
    add_body(doc, "Queda expresamente excluida de este contrato cualquier entrega de código, cesión técnica, auditoría externa o intervención de terceros sobre repositorios, despliegues o infraestructura, salvo autorización expresa y por escrito del desarrollador en un documento separado.")


def add_visuals(doc: Document):
    add_heading(doc, "10. Evidencia visual", 1)
    add_body(doc, "Las siguientes capturas muestran pantallas reales o de preparación de la plataforma web y móvil, incluidas como evidencia visual de alcance y estado de entrega.")
    web = ASSETS_DIR / "screens-web.jpg"
    mobile = ASSETS_DIR / "screens-mobile.jpg"
    if web.exists():
        doc.add_picture(str(web), width=Cm(17.0))
        add_caption(doc, "Figura 4. Pantallas web: login, panel de administración, directorio de usuarios y pagos/avisos.")
    if mobile.exists():
        doc.add_picture(str(mobile), width=Cm(17.0))
        add_caption(doc, "Figura 5. Pantallas móviles iOS/Android adaptadas al uso en teléfono.")


def add_annexes(doc: Document):
    add_heading(doc, "11. Anexos técnicos", 1)
    add_heading(doc, "11.1 Rutas API críticas", 2)
    make_table(doc, ["Área", "Rutas principales", "Uso"], [
        ["Auth", "POST /api/auth/login, POST /api/auth/register", "Login, JWT y alta básica."],
        ["Usuarios", "GET/POST/PUT/DELETE /api/users, PUT /api/users/:id/password", "Directorio, roles, cuotas, perfiles y restablecimiento admin de contraseña."],
        ["Membresías", "POST /api/membership/payments, GET /pending, POST /confirm, POST /users/:id/renew", "Reporte y confirmación de pagos, renovación manual y vencimientos."],
        ["Clases", "GET/POST/PUT/DELETE /api/classes", "Consulta, creación, edición y eliminación de clases."],
        ["Reservas", "POST /reserve, POST /cancel, DELETE /reservations/:userId, POST /attendance/validate", "Reserva, cancelación, promoción de lista y validación de asistencia."],
        ["Push", "POST /api/push-devices", "Registro/actualización de token push tras login móvil."],
        ["Ajustes", "GET/PUT /api/settings, /api/image-bank", "Marca, hero, logos e imágenes de clases."],
        ["Avisos admin", "GET /api/admin-notifications, PATCH /:id/read", "Avisos de pagos reportados y cuotas vencidas."],
    ], [3.0, 6.5, 7.5], font_size=7.6)
    add_heading(doc, "11.2 Estados principales", 2)
    make_table(doc, ["Tipo", "Estados"], [
        ["Roles", "ADMIN, TEACHER, CLIENT"],
        ["Mensualidad", "PAGADO, IMPAGADO"],
        ["Reserva", "CONFIRMADA, EN_ESPERA, ASISTENCIA_VALIDADA, NO_ASISTE, CANCELADA"],
        ["Pago mensualidad", "PENDING_ADMIN_REVIEW, CONFIRMED, REJECTED"],
        ["Plataforma push", "IOS, ANDROID"],
    ], [5.0, 12.0], font_size=8.4)
    doc.add_page_break()
    add_heading(doc, "11.3 Firmas", 2)
    add_body(doc, "En prueba de conformidad, ambas partes firman el presente documento en la fecha indicada.")
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, header in enumerate(["Desarrollador", "Cliente"]):
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=NAVY, size=10)
    rows = [
        (f"Nombre: {DEVELOPER}", "Nombre: ________________________________"),
        ("DNI/NIF: ________________________________", "DNI/NIF/CIF: ____________________________"),
        ("Firma:", "Firma:"),
        ("\n\n\n", "\n\n\n"),
        (f"Fecha: {DOC_DATE}", "Fecha: ____ / ____ / ______"),
    ]
    for row_index, row_data in enumerate(rows, start=1):
        for col_index, value in enumerate(row_data):
            set_cell_text(table.rows[row_index].cells[col_index], value, size=10)


def build_document():
    generate_diagrams_and_collages()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_index(doc)
    add_scope(doc)
    add_deliverables(doc)
    add_roles(doc)
    add_architecture(doc)
    add_data_and_flows(doc)
    add_security_quality(doc)
    add_costs(doc)
    add_maintenance(doc)
    add_contract_terms(doc)
    add_visuals(doc)
    add_annexes(doc)
    props = doc.core_properties
    props.author = DEVELOPER
    props.title = DOC_TITLE
    props.subject = "Contrato y anexo técnico de Ronquillo Te Cuida"
    props.keywords = "Ronquillo Te Cuida, contrato, mantenimiento, Android, iOS, web"
    props.comments = ""
    out = OUT_DIR / "Contrato_Ronquillo_Te_Cuida_Alvaro_Pavon_Martinez_FINAL.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build_document())
